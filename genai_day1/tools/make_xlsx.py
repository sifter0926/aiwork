"""실습용(데이터만) / 강사용(정답: 상관·회귀·차트·보고서) 엑셀 파일 생성."""
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.trendline import Trendline
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, out_practice, out_answer = sys.argv[1:4]
df = pd.read_csv(src)

NAVY, CREAM, ALT = "12304D", "F5F4EF", "ECEEEA"
F = lambda **k: Font(name="Arial", **k)
thin = Side(style="thin", color="C9C9C7")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEADERS = ["날짜", "요일", "주말", "최고기온(℃)", "강수량(mm)", "방문객수(명)", "아이스음료(잔)", "따뜻한음료(잔)", "매출액(천원)"]
WIDTHS = [12, 6, 6, 12, 11, 12, 13, 13, 12]
N = len(df) + 1  # 마지막 데이터 행 (31)


def head(c, text):
    c.value = text
    c.font = F(bold=True, color=CREAM)
    c.fill = HEAD_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BOX


def data_sheet(wb):
    ws = wb.active
    ws.title = "데이터"
    for j, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        head(ws.cell(1, j), h)
        ws.column_dimensions[get_column_letter(j)].width = w
    cols = ["날짜", "요일", "주말", "최고기온", "강수량", "방문객수", "아이스음료", "따뜻한음료", "매출액"]
    for i, row in enumerate(df[cols].itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.font = F()
            c.border = BOX
            c.alignment = Alignment(horizontal="center" if j <= 3 else "right")
            if j == 4:
                c.number_format = "0.0"
            elif j == 9:
                c.number_format = "#,##0"
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ALT)
    ws.freeze_panes = "A2"
    ws.cell(N + 2, 1, "※ 강사 시연용 가상 데이터입니다 (실제 매장 실적 아님). 주말: 토·일=1, 평일=0").font = F(italic=True, color="5E6977")
    return ws


def guide_sheet(wb):
    ws = wb.create_sheet("안내", 0)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 95
    lines = [
        ("엑셀 데이터 분석 보고서 · 실습 파일", F(bold=True, size=16, color=NAVY)),
        ("해운대점 9월 일별 판매 데이터(30일) · 강사 시연용 가상 데이터", F(color="5E6977")),
        ("", None),
        ("① 준비: 파일 > 옵션 > 추가 기능 > [이동] > ‘분석 도구’ 체크 → 데이터 탭에 [데이터 분석] 생성", F()),
        ("② 상관분석: =CORREL(데이터!D2:D31,데이터!G2:G31) / 데이터 분석 > 상관 분석 (입력 범위 D1:I31, 첫째 행 이름표)", F()),
        ("③ 회귀분석: =SLOPE(G2:G31,D2:D31), =INTERCEPT(…), =RSQ(…), =FORECAST(28,G2:G31,D2:D31)", F()),
        ("    데이터 분석 > 회귀 분석: Y축 G1:G31, X축 D1:D31, 이름표 ✓ / 다중회귀: Y축 F1:F31, X축 C1:E31", F()),
        ("④ 시각화: D·G열 선택 > 삽입 > 분산형 > 추세선(선형, 수식·R² 표시) / 콤보 차트(기온은 보조 축)", F()),
        ("⑤ 보고서: 결과 수치를 넣어 AI에게 초안 요청 → 수치·해석 검증 → 1쪽 보고서 완성", F()),
        ("", None),
        ("결과를 적을 칸", F(bold=True, color=NAVY)),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(i, 2, t)
        if f:
            c.font = f
    return ws


# ---------------- 실습용 ----------------
wb = Workbook()
data_sheet(wb)
g = guide_sheet(wb)
fill_rows = [("상관계수 (최고기온–아이스음료)", ""), ("기울기 (SLOPE)", ""), ("절편 (INTERCEPT)", ""), ("결정계수 (RSQ)", ""), ("28℃ 예측 판매량 (FORECAST)", "")]
g.column_dimensions["C"].width = 16
start = 12
for k, (lab, _) in enumerate(fill_rows):
    g.cell(start + k, 2, lab).font = F()
    c = g.cell(start + k, 3)
    c.fill = PatternFill("solid", fgColor="FFFF00")
    c.border = BOX
g.cell(start + len(fill_rows) + 1, 2, "노란 칸에 직접 수식을 넣어 보세요. 예: C12 칸에 =CORREL(데이터!D2:D31,데이터!G2:G31)  → 0.88").font = F(italic=True, color="5E6977")
wb.active = 0
wb.save(out_practice)

# ---------------- 강사용(정답) ----------------
wb = Workbook()
data_sheet(wb)
guide_sheet(wb)
D = "데이터!"
rng = lambda col: f"{D}${col}$2:${col}${N}"
VARS = [("최고기온", "D"), ("강수량", "E"), ("방문객수", "F"), ("아이스음료", "G"), ("따뜻한음료", "H"), ("매출액", "I")]

# 상관분석
ws = wb.create_sheet("상관분석")
ws["A1"] = "상관계수 행렬 (=CORREL) · 색조: 파랑(−1) ↔ 흰색(0) ↔ 주황(+1)"
ws["A1"].font = F(bold=True, size=13, color=NAVY)
head(ws.cell(3, 1), "")
for j, (name, col) in enumerate(VARS, 2):
    head(ws.cell(3, j), name)
    head(ws.cell(3 + j - 1, 1), name)
    ws.column_dimensions[get_column_letter(j)].width = 12
ws.column_dimensions["A"].width = 12
for i, (_, ci) in enumerate(VARS, 4):
    for j, (_, cj) in enumerate(VARS, 2):
        c = ws.cell(i, j, f"=CORREL({rng(ci)},{rng(cj)})")
        c.number_format = "0.00"
        c.font = F()
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
ws.conditional_formatting.add("B4:G9", ColorScaleRule(start_type="num", start_value=-1, start_color="1E4468",
                                                      mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                      end_type="num", end_value=1, end_color="C8692A"))
ws["A11"] = "읽는 법: |r| ≥ 0.7 강함 · 0.3~0.7 보통 · < 0.3 거의 없음 (실무용 어림 기준)"
ws["A11"].font = F(italic=True, color="5E6977")

# 회귀분석
ws = wb.create_sheet("회귀분석")
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 52
ws["A1"] = "단순회귀: 아이스음료(y) ~ 최고기온(x)"
ws["A1"].font = F(bold=True, size=13, color=NAVY)
for j, h in enumerate(["항목", "값", "수식 (y 범위, x 범위 순서)"], 1):
    head(ws.cell(2, j), h)
y, x = rng("G"), rng("D")
ws["B11"] = 28
simple = [
    ("상관계수 r", f"=CORREL({x},{y})", "0.00"),
    ("기울기 b", f"=SLOPE({y},{x})", "0.00"),
    ("절편 a", f"=INTERCEPT({y},{x})", "0.0"),
    ("결정계수 R²", f"=RSQ({y},{x})", "0.00"),
    ("표준 오차", f"=STEYX({y},{x})", "0.0"),
    ("예측: 최고기온 B11℃일 때 아이스음료(잔)", f"=FORECAST(B11,{y},{x})", "0"),
    ("관측한 기온 최솟값(℃)", f"=MIN({x})", "0.0"),
    ("관측한 기온 최댓값(℃)", f"=MAX({x})", "0.0"),
]
for i, (lab, fml, nf) in enumerate(simple, 3):
    ws.cell(i, 1, lab).font = F()
    c = ws.cell(i, 2, fml)
    c.number_format = nf
    c.font = F(bold=True)
    t = ws.cell(i, 3, fml.replace(D, "").replace("$", ""))
    t.data_type = "s"  # 수식을 글자로 보여 주기
    t.font = F(color="5E6977")
    for j in (1, 2, 3):
        ws.cell(i, j).border = BOX
ws["A11"] = "예측할 최고기온 입력(℃) ▶"
ws["A11"].font = F(bold=True)
ws["B11"].fill = PatternFill("solid", fgColor="FFFF00")
ws["B11"].font = F(color="0000FF", bold=True)
ws["C11"] = "노란 칸 값을 바꾸면 B8 예측값이 다시 계산됩니다 (23.5~31.5℃ 안에서만 의미 있음)"
ws["C11"].font = F(italic=True, color="5E6977")

ws["A14"] = "다중회귀: 방문객수(y) ~ 주말 + 최고기온 + 강수량 (LINEST)"
ws["A14"].font = F(bold=True, size=13, color=NAVY)
for j, h in enumerate(["항목", "값", "해석"], 1):
    head(ws.cell(15, j), h)
L = f"LINEST({rng('F')},{D}$C$2:$E${N},TRUE,TRUE)"
multi = [
    ("주말 계수 (명)", f"=INDEX({L},1,3)", "0.0", "주말이면 평일보다 약 87명 더 방문"),
    ("최고기온 계수 (명/℃)", f"=INDEX({L},1,2)", "0.00", "1℃당 약 3명 · 영향 작음"),
    ("강수량 계수 (명/mm)", f"=INDEX({L},1,1)", "0.00", "비 10mm당 약 28명 감소"),
    ("Y 절편", f"=INDEX({L},1,4)", "0.0", ""),
    ("결정계수 R²", f"=INDEX({L},3,1)", "0.00", "방문객 변동의 93%를 세 변수로 설명"),
    ("참고: 기온만 넣었을 때 R²", f"=RSQ({rng('F')},{rng('D')})", "0.00", "기온 하나로는 거의 설명 못 함"),
    ("평일 평균 방문객", f"=AVERAGEIF({rng('C')},0,{rng('F')})", "0.0", "=AVERAGEIF(C2:C31,0,F2:F31)"),
    ("주말 평균 방문객", f"=AVERAGEIF({rng('C')},1,{rng('F')})", "0.0", "=AVERAGEIF(C2:C31,1,F2:F31)"),
]
for i, (lab, fml, nf, note) in enumerate(multi, 16):
    ws.cell(i, 1, lab).font = F()
    c = ws.cell(i, 2, fml)
    c.number_format = nf
    c.font = F(bold=True)
    t = ws.cell(i, 3, note)
    t.data_type = "s"
    t.font = F(color="5E6977")
    for j in (1, 2, 3):
        ws.cell(i, j).border = BOX
ws["A25"] = "LINEST 결과는 X열의 역순(강수량, 최고기온, 주말, 절편)으로 나옵니다. 자세한 P-값은 데이터 분석 > 회귀 분석으로 확인하세요."
ws["A25"].font = F(italic=True, color="5E6977")

# 차트
ws = wb.create_sheet("차트")
ws["A1"] = "시각화 예시 · 산점도+추세선 / 콤보 차트"
ws["A1"].font = F(bold=True, size=13, color=NAVY)
dws = wb["데이터"]
sc = ScatterChart()
sc.title = "최고기온과 아이스음료 판매량"
sc.style = 13
sc.x_axis.title = "최고기온(℃)"
sc.y_axis.title = "아이스음료(잔)"
sc.x_axis.scaling.min, sc.x_axis.scaling.max = 22, 33
sc.y_axis.scaling.min, sc.y_axis.scaling.max = 100, 300
sc.x_axis.delete = False
sc.y_axis.delete = False
ser = Series(Reference(dws, min_col=7, min_row=1, max_row=N), Reference(dws, min_col=4, min_row=2, max_row=N), title_from_data=True)
ser.marker.symbol = "circle"
ser.marker.size = 7
ser.marker.graphicalProperties.solidFill = "1F6F78"
ser.marker.graphicalProperties.line.solidFill = "1F6F78"
ser.graphicalProperties.line.noFill = True
ser.trendline = Trendline(trendlineType="linear", dispEq=True, dispRSqr=True)
sc.series.append(ser)
sc.legend = None
sc.width, sc.height = 18, 10
ws.add_chart(sc, "A3")

bar = BarChart()
bar.title = "9월 일별 아이스음료 판매량과 최고기온"
bar.add_data(Reference(dws, min_col=7, min_row=1, max_row=N), titles_from_data=True)
bar.set_categories(Reference(dws, min_col=1, min_row=2, max_row=N))
bar.y_axis.title = "잔"
bar.x_axis.delete = False
bar.y_axis.delete = False
line = LineChart()
line.add_data(Reference(dws, min_col=4, min_row=1, max_row=N), titles_from_data=True)
line.y_axis.axId = 200
line.y_axis.title = "℃"
line.y_axis.crosses = "max"
line.y_axis.scaling.min, line.y_axis.scaling.max = 15, 35
line.y_axis.delete = False
line.y_axis.majorGridlines = None
bar += line
bar.legend.position = "b"
bar.width, bar.height = 24, 10
ws.add_chart(bar, "A25")

# 보고서 (수치는 모두 수식으로 연결)
ws = wb.create_sheet("보고서")
ws.column_dimensions["A"].width = 120
R = "회귀분석!"
rows = [
    ("해운대점 9월 판매 데이터 분석 보고 (운영지원팀)", F(bold=True, size=14, color=NAVY)),
    ("[요약]", F(bold=True)),
    (f'="① 아이스음료 판매는 최고기온과 강한 양의 상관(r="&TEXT({R}B3,"0.00")&")이 있고, 1℃ 오를 때 약 "&TEXT({R}B4,"0")&"잔 늘어나는 경향"', F()),
    (f'="② 방문객 수는 주말(+"&TEXT({R}B16,"0")&"명)과 강수량(1mm당 "&TEXT({R}B18,"0.0")&"명)의 영향이 크고 기온 영향은 작음"', F()),
    ('="③ 방문객 수와 매출은 거의 같이 움직임(r="&TEXT(상관분석!G6,"0.00")&")"', F()),
    ("[핵심 지표]", F(bold=True)),
    (f'="회귀식: 아이스음료 = "&TEXT({R}B4,"0.00")&" × 최고기온 − "&TEXT(-{R}B5,"0.0")&" · R² "&TEXT({R}B6,"0.00")&" · "&{R}B11&"℃ 예측 "&TEXT({R}B8,"0")&"잔"', F()),
    (f'="방문객 다중회귀 R² "&TEXT({R}B20,"0.00")&" (기온만 넣으면 "&TEXT({R}B21,"0.00")&")"', F()),
    ("[시사점]", F(bold=True)),
    (f'="① 다음 날 최고기온 예보로 아이스음료 재료 준비량 조정 ("&{R}B11&"℃ ≈ "&TEXT({R}B8,"0")&"잔)"', F()),
    (f'="② 주말은 평일보다 방문객이 평균 "&TEXT({R}B23-{R}B22,"0")&"명 많아("&TEXT({R}B22,"0")&"→"&TEXT({R}B23,"0")&"명) 인력 보강 검토"', F()),
    ("③ 비 예보일은 방문객 감소에 대비해 포장·배달 프로모션 검토", F()),
    ("[한계]", F(bold=True)),
    (f'="9월 30일 자료(최고기온 "&TEXT({R}B9,"0.0")&"~"&TEXT({R}B10,"0.0")&"℃)로 분석해 다른 계절·지점에 그대로 적용할 수 없으며, 상관관계는 원인을 뜻하지 않음"', F()),
]
for i, (t, f) in enumerate(rows, 1):
    c = ws.cell(i, 1, t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
wb.active = 0
wb.save(out_answer)
print("saved")
